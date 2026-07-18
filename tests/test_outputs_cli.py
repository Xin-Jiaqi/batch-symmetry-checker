from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from pymatgen.core import Lattice, Structure

from batch_symmetry_checker import AnalysisConfig, analyze_directory
from batch_symmetry_checker.cli import EXIT_INCOMPLETE, EXIT_SUCCESS, EXIT_USAGE, main
from batch_symmetry_checker.models import REPORT_SCHEMA_VERSION
from batch_symmetry_checker.outputs import ERROR_FIELDS, RECORD_FIELDS, write_report


def write_poscar(path: Path) -> None:
    Structure(Lattice.cubic(4.0), ["Si"], [[0, 0, 0]]).to(filename=str(path), fmt="poscar")


def test_json_report_has_versioned_complete_contract(tmp_path: Path) -> None:
    write_poscar(tmp_path / "POSCAR")
    output = tmp_path / "report.json"
    assert (
        main(["--input", str(tmp_path), "--output", str(output), "--tolerances", "1e-3", "--quiet"])
        == EXIT_SUCCESS
    )
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["schema_version"] == REPORT_SCHEMA_VERSION
    assert payload["status"] == "complete"
    assert payload["summary"] == {"record_count": 1, "error_count": 0, "source_count": 1}
    assert payload["records"][0]["source_sha256"]
    assert payload["errors"] == []


def test_csv_and_error_sidecar_share_machine_readable_fields(tmp_path: Path) -> None:
    write_poscar(tmp_path / "POSCAR")
    (tmp_path / "POSCAR_bad").write_text("broken\n", encoding="utf-8")
    output = tmp_path / "report.csv"
    assert (
        main(["--input", str(tmp_path), "--output", str(output), "--tolerances", "1e-3", "--quiet"])
        == EXIT_INCOMPLETE
    )
    with output.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    with (tmp_path / "report.errors.csv").open(encoding="utf-8", newline="") as handle:
        errors = list(csv.DictReader(handle))
    assert rows[0]["source_path"] == "POSCAR"
    assert errors[0]["source_path"] == "POSCAR_bad"
    assert errors[0]["code"] == "structure_parse_error"


def test_excel_has_records_errors_and_metadata_sheets(tmp_path: Path) -> None:
    import openpyxl

    write_poscar(tmp_path / "POSCAR")
    output = tmp_path / "report.xlsx"
    assert (
        main(["--input", str(tmp_path), "--output", str(output), "--tolerances", "1e-3", "--quiet"])
        == EXIT_SUCCESS
    )
    workbook = openpyxl.load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["symmetry_records", "errors", "metadata"]


def test_missing_input_and_bad_tolerance_are_usage_errors(tmp_path: Path) -> None:
    assert main(["--input", str(tmp_path / "missing"), "--quiet"]) == EXIT_USAGE
    assert main(["--input", str(tmp_path), "--tolerances", "0", "--quiet"]) == EXIT_USAGE


def test_no_files_is_usage_error(tmp_path: Path) -> None:
    assert main(["--input", str(tmp_path), "--quiet"]) == EXIT_USAGE


def test_existing_output_requires_force_and_is_not_modified(tmp_path: Path) -> None:
    write_poscar(tmp_path / "POSCAR")
    output = tmp_path / "report.json"
    output.write_text("sentinel", encoding="utf-8")
    args = ["--input", str(tmp_path), "--output", str(output), "--tolerances", "1e-3", "--quiet"]
    assert main(args) == EXIT_USAGE
    assert output.read_text(encoding="utf-8") == "sentinel"
    assert main([*args, "--force"]) == EXIT_SUCCESS
    assert json.loads(output.read_text(encoding="utf-8"))["status"] == "complete"


def test_csv_error_sidecar_collision_is_checked_before_any_write(tmp_path: Path) -> None:
    write_poscar(tmp_path / "POSCAR")
    (tmp_path / "POSCAR_bad").write_text("broken\n", encoding="utf-8")
    output = tmp_path / "report.csv"
    error_output = tmp_path / "report.errors.csv"
    error_output.write_text("sentinel", encoding="utf-8")
    args = ["--input", str(tmp_path), "--output", str(output), "--tolerances", "1e-3", "--quiet"]
    assert main(args) == EXIT_USAGE
    assert not output.exists()
    assert error_output.read_text(encoding="utf-8") == "sentinel"
    assert main([*args, "--force"]) == EXIT_INCOMPLETE
    assert output.exists()
    assert "structure_parse_error" in error_output.read_text(encoding="utf-8")


def test_empty_csv_and_xlsx_tables_keep_fixed_schemas(tmp_path: Path) -> None:
    import openpyxl

    inputs = tmp_path / "inputs"
    inputs.mkdir()
    (inputs / "POSCAR_bad").write_text("broken\n", encoding="utf-8")

    csv_output = tmp_path / "empty.csv"
    assert main(["--input", str(inputs), "--output", str(csv_output), "--quiet"]) == EXIT_INCOMPLETE
    with csv_output.open(encoding="utf-8", newline="") as handle:
        assert csv.DictReader(handle).fieldnames == RECORD_FIELDS
    with (tmp_path / "empty.errors.csv").open(encoding="utf-8", newline="") as handle:
        assert csv.DictReader(handle).fieldnames == ERROR_FIELDS

    xlsx_output = tmp_path / "empty.xlsx"
    assert (
        main(["--input", str(inputs), "--output", str(xlsx_output), "--quiet"])
        == EXIT_INCOMPLETE
    )
    workbook = openpyxl.load_workbook(xlsx_output, read_only=True)
    assert [cell.value for cell in next(workbook["symmetry_records"].iter_rows())] == RECORD_FIELDS
    assert [cell.value for cell in next(workbook["errors"].iter_rows())] == ERROR_FIELDS


def test_successful_csv_rerun_replaces_stale_errors_with_empty_sidecar(tmp_path: Path) -> None:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    write_poscar(inputs / "POSCAR")
    broken = inputs / "POSCAR_bad"
    broken.write_text("broken\n", encoding="utf-8")
    output = tmp_path / "report.csv"
    args = ["--input", str(inputs), "--output", str(output), "--quiet"]
    assert main(args) == EXIT_INCOMPLETE
    broken.unlink()
    assert main([*args, "--force"]) == EXIT_SUCCESS
    with (tmp_path / "report.errors.csv").open(encoding="utf-8", newline="") as handle:
        assert list(csv.DictReader(handle)) == []


def test_failed_csv_write_preserves_existing_report_pair(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    write_poscar(inputs / "POSCAR")
    report = analyze_directory(AnalysisConfig(inputs), producer_version="test")
    output = tmp_path / "report.csv"
    error_output = tmp_path / "report.errors.csv"
    output.write_text("old-main\n", encoding="utf-8")
    error_output.write_text("old-errors\n", encoding="utf-8")

    from batch_symmetry_checker import outputs

    original = outputs._write_csv_rows
    calls = 0

    def fail_on_sidecar(path: Path, rows: list[dict[str, object]], names: list[str]) -> None:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("simulated sidecar failure")
        original(path, rows, names)

    monkeypatch.setattr(outputs, "_write_csv_rows", fail_on_sidecar)
    with pytest.raises(OSError, match="simulated"):
        write_report(report, output, overwrite=True)
    assert output.read_text(encoding="utf-8") == "old-main\n"
    assert error_output.read_text(encoding="utf-8") == "old-errors\n"


def test_failed_main_csv_replace_rolls_back_sidecar(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    write_poscar(inputs / "POSCAR")
    report = analyze_directory(AnalysisConfig(inputs), producer_version="test")
    output = tmp_path / "report.csv"
    error_output = tmp_path / "report.errors.csv"
    output.write_text("old-main\n", encoding="utf-8")
    error_output.write_text("old-errors\n", encoding="utf-8")

    from batch_symmetry_checker import outputs

    original = outputs._replace_temporary
    calls = 0

    def fail_on_main(temporary: Path, destination: Path) -> None:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("simulated main replacement failure")
        original(temporary, destination)

    monkeypatch.setattr(outputs, "_replace_temporary", fail_on_main)
    with pytest.raises(OSError, match="simulated"):
        write_report(report, output, overwrite=True)
    assert output.read_text(encoding="utf-8") == "old-main\n"
    assert error_output.read_text(encoding="utf-8") == "old-errors\n"


def test_failed_sidecar_replace_restores_previous_sidecar(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    write_poscar(inputs / "POSCAR")
    report = analyze_directory(AnalysisConfig(inputs), producer_version="test")
    output = tmp_path / "report.csv"
    error_output = tmp_path / "report.errors.csv"
    output.write_text("old-main\n", encoding="utf-8")
    error_output.write_text("old-errors\n", encoding="utf-8")

    from batch_symmetry_checker import outputs

    def fail_on_sidecar(temporary: Path, destination: Path) -> None:
        raise OSError("simulated sidecar replacement failure")

    monkeypatch.setattr(outputs, "_replace_temporary", fail_on_sidecar)
    with pytest.raises(OSError, match="simulated"):
        write_report(report, output, overwrite=True)
    assert output.read_text(encoding="utf-8") == "old-main\n"
    assert error_output.read_text(encoding="utf-8") == "old-errors\n"
